from datetime import date
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = 'MANUAL_BOOK_WEBSITE_GODONG_IJO.docx'

def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    properties.append(repeat)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        set_cell_text(header.cells[index], text, bold=True, color=(255, 255, 255))
        set_cell_shading(header.cells[index], '1F4E5F')
        set_cell_margins(header.cells[index])
    for row_data in rows:
        row = table.add_row()
        for index, value in enumerate(row_data):
            set_cell_text(row.cells[index], value)
            set_cell_margins(row.cells[index])
            if len(table.rows) % 2 == 0:
                set_cell_shading(row.cells[index], 'EAF3F5')
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return table

def add_bullets(doc, items, level=0):
    for item in items:
        paragraph = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
        paragraph.add_run(item)

def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style='List Number')
        paragraph.add_run(item)

def add_note(doc, title, text, color='FFF2CC'):
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title + ': ')
    run.bold = True
    paragraph.add_run(text)
    doc.add_paragraph()

def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.style = 'No Spacing'
    run = paragraph.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(40, 40, 40)
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(4)

def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)

def add_para(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph

def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [('Title', 28, '1F4E5F'), ('Heading 1', 18, '1F4E5F'), ('Heading 2', 14, '287D8B'), ('Heading 3', 11, '1F4E5F')]:
        style = styles[name]
        style.font.name = 'Aptos Display'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    for style_name in ['List Bullet', 'List Bullet 2', 'List Number']:
        styles[style_name].font.name = 'Aptos'
        styles[style_name].font.size = Pt(10.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Manual Book Website Godong Ijo  |  Dokumen operasional').font.size = Pt(8)

def add_cover(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(70)
    run = paragraph.add_run('GODONG IJO')
    run.font.name = 'Aptos Display'
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 95)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('MANUAL BOOK WEBSITE')
    run.font.size = Pt(23)
    run.bold = True
    run.font.color.rgb = RGBColor(40, 125, 139)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    run = paragraph.add_run('Panduan penggunaan, pengelolaan booking, pembayaran, dan administrasi sistem')
    run.font.size = Pt(13)
    run.italic = True

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    data = [
        ('Nama sistem', 'Website Godong Ijo'),
        ('Jenis aplikasi', 'Website wisata, paket rekreasi, booking, dan panel admin'),
        ('Versi dokumen', '1.0'),
        ('Tanggal dokumen', '8 September 2026'),
    ]
    for row, values in zip(table.rows, data):
        for index, value in enumerate(values):
            set_cell_text(row.cells[index], value, bold=index == 0)
            set_cell_margins(row.cells[index])
            if index == 0:
                set_cell_shading(row.cells[index], 'EAF3F5')
    doc.add_paragraph()
    add_note(doc, 'Tujuan dokumen', 'Manual ini menjelaskan cara website bekerja dari sudut pandang pengunjung, pelanggan, dan admin. Gunakan bagian sesuai peran pengguna.')
    doc.add_page_break()

def build_document():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, 'Daftar Isi', 1)
    toc_items = [
        '1. Gambaran Umum Sistem',
        '2. Peran Pengguna dan Hak Akses',
        '3. Panduan Pengunjung dan Pelanggan',
        '4. Panduan Booking Reguler',
        '5. Panduan Booking Fishing',
        '6. Pembayaran Midtrans',
        '7. Konfirmasi Booking dan E-ticket',
        '8. Panduan Login dan Panel Admin',
        '9. Manajemen Booking',
        '10. Manajemen Paket Wisata',
        '11. Customer, Laporan, Aktivitas, dan Notifikasi',
        '12. Pengaturan Sistem dan Email',
        '13. Chatbot dan Fitur Pendukung',
        '14. Alur Status Booking',
        '15. Troubleshooting',
        '16. Keamanan dan Praktik Operasional',
        '17. Checklist Operasional Admin',
        '18. Ringkasan Endpoint dan Komponen Sistem',
    ]
    add_numbered(doc, toc_items)
    add_note(doc, 'Catatan navigasi', 'Judul bab dibuat sebagai heading Word sehingga dapat dipakai untuk Navigation Pane. Jika diperlukan, daftar isi otomatis dapat dibuat melalui References > Table of Contents di Microsoft Word.')

    add_heading(doc, '1. Gambaran Umum Sistem', 1)
    add_para(doc, 'Website Godong Ijo adalah platform informasi wisata dan kuliner yang menyediakan halaman destinasi, paket wisata, wisata edukasi, kontak, booking online, pembayaran Midtrans, notifikasi, dan panel administrasi.')
    add_para(doc, 'Pengunjung dapat melihat informasi tanpa login. Booking dilakukan sebagai guest menggunakan nama, email, nomor WhatsApp, tanggal kunjungan, dan jumlah peserta. Admin mengelola booking serta konten paket melalui panel admin yang terlindungi login.')
    add_table(doc, ['Area sistem', 'Fungsi utama', 'Pengguna'], [
        ('Website publik', 'Informasi destinasi, paket, edukasi, kontak, chatbot, dan CTA booking', 'Pengunjung'),
        ('Booking dan pembayaran', 'Validasi data, cek kuota, membuat order, pembayaran, status, dan e-ticket', 'Pelanggan dan admin'),
        ('Panel admin', 'Dashboard, booking, paket, customer, laporan, notifikasi, setting', 'Admin/operator'),
        ('Integrasi', 'Midtrans, SMTP email, WhatsApp, PDF e-ticket, scheduler', 'Sistem dan admin'),
    ], widths=[1.45, 3.65, 1.45])

    add_heading(doc, '2. Peran Pengguna dan Hak Akses', 1)
    add_heading(doc, '2.1 Pengunjung atau pelanggan guest', 2)
    add_bullets(doc, [
        'Melihat halaman utama, destinasi, kategori paket, detail paket, wisata edukasi, sekolah mitra, dan kontak.',
        'Mengirim pertanyaan melalui chatbot publik.',
        'Membuat booking tanpa akun pelanggan.',
        'Membayar melalui Midtrans sesuai order yang dibuat.',
        'Melihat konfirmasi dan mengunduh e-ticket menggunakan kode booking.',
    ])
    add_heading(doc, '2.2 Admin atau operator', 2)
    add_bullets(doc, [
        'Login melalui /admin/login.',
        'Mengelola booking, status pembayaran, email, dan export data.',
        'Mengelola paket wisata, harga, kuota, status aktif, dan foto.',
        'Melihat data customer, laporan, aktivitas, notifikasi, dan chatbot admin.',
        'Mengubah pengaturan admin, kontak, notifikasi, dan email.',
    ])
    add_note(doc, 'Penting', 'Semua route admin dilindungi middleware admin. Hak akses saat ini belum dibagi berdasarkan role, sehingga admin yang berhasil login memiliki akses ke seluruh menu admin.')

    add_heading(doc, '3. Panduan Pengunjung dan Pelanggan', 1)
    add_heading(doc, '3.1 Halaman utama', 2)
    add_para(doc, 'Halaman utama menampilkan hero section, destinasi unggulan, statistik, paket aktif, informasi fasilitas, CTA booking, informasi kontak, WhatsApp, dan chatbot.')
    add_heading(doc, '3.2 Halaman destinasi', 2)
    add_table(doc, ['URL', 'Isi halaman'], [
        ('/destinasi/the-waterfall', 'Informasi The Waterfall Resto, fasilitas, galeri, jam buka, dan CTA booking.'),
        ('/destinasi/monster-fish', 'Informasi Monster Fish Fishing Lake, aktivitas sport fishing, fasilitas, dan aturan.'),
    ], widths=[2.15, 4.4])
    add_heading(doc, '3.3 Halaman kategori dan detail paket', 2)
    add_bullets(doc, [
        'Kategori tersedia: The Waterfall Resto, Private Room, dan Fishing Lake.',
        'Kartu paket menampilkan nama, foto, harga, status aktif, dan tombol booking.',
        'Halaman detail menampilkan deskripsi, fasilitas, galeri, itinerary, informasi khusus, dan CTA booking.',
    ])
    add_heading(doc, '3.4 Wisata edukasi dan kontak', 2)
    add_para(doc, 'Halaman Wisata Edukasi menampilkan program Virtual Fieldtrip, Goes To School, Fieldtrip at Godong Ijo, edukasi lingkungan, sains, dan seni. Reservasi program edukasi diarahkan melalui WhatsApp, bukan checkout otomatis.')
    add_para(doc, 'Halaman Kontak menampilkan alamat, jam operasional, nomor telepon, email, peta, dan tautan WhatsApp.')

    add_heading(doc, '4. Panduan Booking Reguler', 1)
    add_heading(doc, '4.1 Cara membuat booking', 2)
    add_numbered(doc, [
        'Buka website dan pilih paket yang diinginkan.',
        'Klik tombol Booking atau CTA terkait paket.',
        'Isi nama lengkap, email, nomor WhatsApp, tanggal kunjungan, dan jumlah orang.',
        'Untuk paket Private Room, pastikan jumlah peserta memenuhi minimum paket.',
        'Periksa kembali data dan kirim formulir.',
        'Sistem memvalidasi paket aktif, harga, tanggal, nomor telepon, jumlah peserta, dan kuota.',
        'Jika valid, sistem membuat booking dengan status pending dan mengurangi kuota jadwal secara aman.',
        'Sistem membuat order pembayaran atau token simulasi sesuai konfigurasi.',
        'Lanjutkan pembayaran dan simpan kode booking.',
    ])
    add_heading(doc, '4.2 Validasi yang berlaku', 2)
    add_table(doc, ['Data', 'Aturan umum'], [
        ('Nama', 'Wajib, teks, maksimal 255 karakter.'),
        ('Email', 'Wajib dan harus berformat email.'),
        ('Nomor WhatsApp', 'Wajib, format diawali 08 atau 62, dan panjang sesuai aturan validasi.'),
        ('Tanggal kunjungan', 'Wajib, tidak boleh sebelum hari ini.'),
        ('Jumlah orang', 'Wajib, minimal 1 dan dibatasi maksimal 100 pada booking reguler.'),
        ('Kuota', 'Harus mencukupi untuk tanggal yang dipilih.'),
    ], widths=[1.45, 5.1])
    add_note(doc, 'Jika gagal', 'Jangan mengirim formulir berulang kali tanpa membaca pesan error. Periksa tanggal, nomor WhatsApp, jumlah peserta, dan status paket terlebih dahulu.')

    add_heading(doc, '5. Panduan Booking Fishing', 1)
    add_para(doc, 'Booking fishing memakai formulir khusus karena memiliki jenis pemancingan, jumlah joran, jam kunjungan, durasi, tambahan alat, dan aturan persetujuan.')
    add_heading(doc, '5.1 Data yang perlu diisi', 2)
    add_bullets(doc, [
        'Nama lengkap, email, nomor WhatsApp, dan tanggal kunjungan.',
        'Jam kunjungan antara 09:00 sampai 21:00.',
        'Jenis pemancingan: tarikan, sewa joran, jackpot, atau kiloan.',
        'Jumlah joran, maksimal 999 sesuai validasi aplikasi.',
        'Durasi dan tambahan jam untuk jenis tarikan.',
        'Ukuran joran untuk jenis sewa joran.',
        'Pilihan sewa alat, anak ikan komet, dan umpan jadi jika diperlukan.',
        'Persetujuan terhadap aturan pemancingan.',
    ])
    add_heading(doc, '5.2 Perhitungan harga', 2)
    add_table(doc, ['Komponen', 'Keterangan'], [
        ('Tarikan atau jenis berharga', 'Harga dasar dikalikan jumlah joran.'),
        ('Tambahan jam', 'Rp20.000 per jam tambahan.'),
        ('Sewa alat', 'Rp50.000 jika dipilih.'),
        ('Anak ikan komet', 'Rp15.000 per unit.'),
        ('Umpan jadi', 'Rp25.000 per unit.'),
        ('Kiloan', 'Estimasi awal dapat kosong; total dihitung saat penimbangan.'),
    ], widths=[1.7, 4.85])
    add_note(doc, 'Kuota fishing', 'Sistem mengunci jadwal dalam transaksi dan menolak booking jika kuota joran pada tanggal tersebut tidak mencukupi.')

    add_heading(doc, '6. Pembayaran Midtrans', 1)
    add_heading(doc, '6.1 Cara kerja pembayaran', 2)
    add_numbered(doc, [
        'Sistem membuat order ID dan nominal pembayaran setelah booking valid.',
        'Sistem meminta Snap Token ke Midtrans, kecuali mode simulation aktif.',
        'Pelanggan menyelesaikan pembayaran pada halaman atau popup Midtrans.',
        'Midtrans mengirim notifikasi ke endpoint /midtrans/notification.',
        'Sistem memeriksa signature SHA-512 dan mencocokkan gross amount.',
        'Status pembayaran dipetakan ke status internal.',
        'Jika sukses, status pembayaran menjadi success dan status booking menjadi paid.',
        'Email pembayaran berhasil dan notifikasi admin dikirim bila setting terkait aktif.',
    ])
    add_heading(doc, '6.2 Status pembayaran', 2)
    add_table(doc, ['Status Midtrans', 'Status internal', 'Dampak'], [
        ('settlement', 'success', 'Booking menjadi paid.'),
        ('capture dengan fraud accept', 'success', 'Booking menjadi paid.'),
        ('capture dengan fraud deny', 'failed', 'Pembayaran gagal.'),
        ('pending atau challenge', 'pending', 'Menunggu keputusan pembayaran.'),
        ('cancel atau deny', 'failed', 'Pembayaran gagal.'),
        ('expire', 'expired', 'Pembayaran kadaluarsa.'),
    ], widths=[2.0, 1.6, 2.95])
    add_heading(doc, '6.3 Mode pembayaran', 2)
    add_bullets(doc, [
        'simulation: dipakai untuk pengujian lokal tanpa request pembayaran nyata.',
        'live: dipakai untuk integrasi Midtrans sesuai konfigurasi server key dan mode produksi/sandbox.',
    ])
    add_note(doc, 'Keamanan key', 'Jangan menulis Server Key, Client Key, password SMTP, atau APP_KEY ke dalam manual, source code, screenshot, atau chat publik. Simpan hanya pada file environment server.')

    add_heading(doc, '7. Konfirmasi Booking dan E-ticket', 1)
    add_heading(doc, '7.1 Halaman konfirmasi', 2)
    add_para(doc, 'Halaman konfirmasi menampilkan kode booking, pelanggan, paket, tanggal kunjungan, jumlah peserta atau joran, status booking, status pembayaran, serta tautan bantuan WhatsApp.')
    add_heading(doc, '7.2 Syarat e-ticket', 2)
    add_bullets(doc, [
        'Booking berstatus paid.',
        'Pembayaran berstatus success.',
        'Data booking dan relasi paket tersedia.',
        'DomPDF berhasil digunakan oleh aplikasi.',
    ])
    add_numbered(doc, [
        'Buka halaman konfirmasi menggunakan kode booking.',
        'Pastikan status pembayaran sudah berhasil.',
        'Klik tombol Download E-ticket.',
        'Simpan PDF dan tunjukkan saat diperlukan di lokasi.',
    ])

    add_heading(doc, '8. Panduan Login dan Panel Admin', 1)
    add_heading(doc, '8.1 Login admin', 2)
    add_numbered(doc, [
        'Buka /admin/login.',
        'Masukkan email atau username dan password admin.',
        'Gunakan Remember Me hanya pada perangkat yang aman.',
        'Setelah berhasil login, sistem mengarahkan ke dashboard.',
        'Gunakan Logout setelah selesai, terutama pada komputer bersama.',
    ])
    add_note(doc, 'Instalasi awal', 'Jika aplikasi masih memakai akun admin default dari seeder, segera ganti password melalui menu Settings sebelum sistem digunakan di lingkungan nyata.')
    add_heading(doc, '8.2 Dashboard admin', 2)
    add_bullets(doc, [
        'Booking hari ini dan bulan ini.',
        'Pendapatan dari booking paid.',
        'Jumlah booking pending.',
        'Lima booking terbaru.',
        'Navigasi cepat ke menu operasional.',
    ])

    add_heading(doc, '9. Manajemen Booking', 1)
    add_heading(doc, '9.1 Daftar booking', 2)
    add_para(doc, 'Menu Bookings menyediakan pagination, pencarian kode booking/nama/email, filter status, dan filter rentang tanggal kunjungan.')
    add_heading(doc, '9.2 Detail dan perubahan status', 2)
    add_table(doc, ['Status', 'Kapan digunakan', 'Catatan'], [
        ('pending', 'Booking dibuat tetapi belum lunas.', 'Pantau pembayaran atau hubungi pelanggan.'),
        ('paid', 'Pembayaran berhasil diverifikasi.', 'Booking siap diproses sesuai SOP lokasi.'),
        ('cancelled', 'Booking dibatalkan.', 'Kuota dapat dikembalikan sesuai aturan.'),
        ('expired', 'Pembayaran atau booking kadaluarsa.', 'Periksa kebutuhan pengembalian kuota.'),
    ], widths=[1.1, 2.55, 2.9])
    add_heading(doc, '9.3 Mengirim email manual', 2)
    add_numbered(doc, [
        'Buka detail booking.',
        'Pilih jenis email: booking, payment, atau cancellation.',
        'Isi alasan bila diperlukan.',
        'Kirim dan periksa pesan hasil pengiriman.',
        'Jika gagal, periksa konfigurasi SMTP dan log aplikasi.',
    ])
    add_heading(doc, '9.4 Export data', 2)
    add_para(doc, 'Gunakan menu export untuk membuat CSV, Excel-compatible XLS, atau tampilan PDF/HTML sesuai kebutuhan laporan.')

    add_heading(doc, '10. Manajemen Paket Wisata', 1)
    add_heading(doc, '10.1 Membuat paket', 2)
    add_numbered(doc, [
        'Buka Admin > Packages > Tambah Paket.',
        'Isi nama, jenis paket, deskripsi, harga, dan kuota.',
        'Tentukan status aktif.',
        'Upload foto dengan format JPEG, PNG, JPG, atau WEBP maksimal 2 MB.',
        'Simpan dan periksa tampilan paket di halaman publik.',
    ])
    add_heading(doc, '10.2 Edit dan nonaktifkan paket', 2)
    add_bullets(doc, [
        'Edit harga atau kuota hanya setelah memastikan tidak mengganggu booking yang sudah ada.',
        'Nonaktifkan paket bila tidak ingin menerima booking baru.',
        'Paket yang memiliki histori booking tidak boleh dihapus; gunakan status nonaktif.',
        'Saat mengganti foto, periksa file baru dan tampilan halaman detail.',
    ])
    add_note(doc, 'Konten publik', 'Setelah mengubah paket, bersihkan cache bila perubahan belum terlihat dan buka ulang halaman dalam mode private/incognito untuk menghindari cache browser.')

    add_heading(doc, '11. Customer, Laporan, Aktivitas, dan Notifikasi', 1)
    add_heading(doc, '11.1 Customer', 2)
    add_para(doc, 'Menu Customers mengelompokkan pelanggan guest berdasarkan email atau nomor telepon. Admin dapat melihat jumlah booking, total belanja, booking pending, dan booking paid.')
    add_heading(doc, '11.2 Reports', 2)
    add_bullets(doc, [
        'Total pendapatan.',
        'Rata-rata nilai booking.',
        'Jumlah booking paid.',
        'Distribusi status booking.',
        'Paket terlaris.',
        'Pendapatan dan jumlah booking harian.',
    ])
    add_heading(doc, '11.3 Activity dan Notifications', 2)
    add_para(doc, 'Activity menampilkan aktivitas operasional terbaru. Notifications menampilkan booking baru dan pembayaran berhasil, termasuk jumlah unread dan aksi tandai sudah dibaca.')

    add_heading(doc, '12. Pengaturan Sistem dan Email', 1)
    add_heading(doc, '12.1 Pengaturan admin', 2)
    add_bullets(doc, [
        'Nama dan email admin.',
        'Password admin.',
        'Nama website.',
        'Email dan nomor kontak.',
        'Aktif/nonaktif notifikasi booking.',
        'Aktif/nonaktif notifikasi email.',
    ])
    add_heading(doc, '12.2 Konfigurasi email', 2)
    add_table(doc, ['Konfigurasi', 'Fungsi'], [
        ('MAIL_MAILER', 'Driver pengiriman email.'),
        ('MAIL_HOST dan MAIL_PORT', 'Server dan port SMTP.'),
        ('MAIL_USERNAME dan MAIL_PASSWORD', 'Kredensial SMTP.'),
        ('MAIL_ENCRYPTION', 'TLS/SSL sesuai provider.'),
        ('MAIL_FROM_ADDRESS dan MAIL_FROM_NAME', 'Identitas pengirim.'),
        ('MAIL_DAILY_LIMIT', 'Batas pengiriman email harian.'),
    ], widths=[2.15, 4.4])
    add_heading(doc, '12.3 Email otomatis dan scheduler', 2)
    add_bullets(doc, [
        'Konfirmasi booking.',
        'Pembayaran berhasil.',
        'Pembatalan.',
        'Perubahan status oleh admin.',
        'Pengingat H-1 pada pukul 08.00 WIB.',
        'Permintaan review pada pukul 09.00 WIB.',
    ])
    add_para(doc, 'Scheduler membutuhkan Task Scheduler Windows atau cron yang menjalankan php artisan schedule:run secara rutin.')

    add_heading(doc, '13. Chatbot dan Fitur Pendukung', 1)
    add_heading(doc, '13.1 Chatbot publik', 2)
    add_bullets(doc, [
        'Menjawab harga paket, jam buka, lokasi, dan cara booking.',
        'Memberi informasi Private Room dan pertanyaan umum.',
        'Mengarahkan ke WhatsApp bila pertanyaan tidak dikenali.',
        'Dibatasi 30 request per menit.',
    ])
    add_heading(doc, '13.2 Chatbot admin', 2)
    add_bullets(doc, [
        'Mencari booking berdasarkan kode, nama, email, atau nomor WhatsApp.',
        'Melihat booking pending dan ringkasan pendapatan.',
        'Menampilkan paket dengan kuota rendah.',
        'Menampilkan jumlah customer dan notifikasi.',
    ])
    add_heading(doc, '13.3 WhatsApp dan sitemap', 2)
    add_para(doc, 'Tautan WhatsApp dipakai untuk pertanyaan dan reservasi edukasi. Sitemap tersedia pada /sitemap.xml untuk membantu mesin pencari menemukan halaman publik.')

    add_heading(doc, '14. Alur Status Booking', 1)
    add_table(doc, ['Tahap', 'Status booking', 'Status pembayaran', 'Aksi sistem'], [
        ('Form valid dikirim', 'pending', 'pending', 'Booking, jadwal, kuota, dan pembayaran dibuat.'),
        ('Menunggu pembayaran', 'pending', 'pending', 'Pelanggan membayar melalui Midtrans.'),
        ('Pembayaran sukses', 'paid', 'success', 'Email sukses, notifikasi admin, e-ticket tersedia.'),
        ('Pembayaran gagal', 'pending atau sesuai keputusan admin', 'failed', 'Admin dapat menghubungi pelanggan atau membatalkan.'),
        ('Dibatalkan', 'cancelled', 'sesuai histori', 'Kuota diproses sesuai aturan pembatalan.'),
        ('Kadaluarsa', 'expired', 'expired', 'Periksa kuota dan kebutuhan tindak lanjut.'),
    ], widths=[1.25, 1.3, 1.45, 3.0])
    add_note(doc, 'Prinsip penting', 'Jangan menandai booking paid hanya berdasarkan screenshot pelanggan. Gunakan status dari Midtrans atau verifikasi transaksi melalui kanal resmi.')

    add_heading(doc, '15. Troubleshooting', 1)
    add_heading(doc, '15.1 Booking gagal dibuat', 2)
    add_numbered(doc, [
        'Pastikan paket masih aktif.',
        'Pastikan harga dan kuota sudah diisi.',
        'Periksa tanggal kunjungan dan jumlah peserta.',
        'Untuk Private Room, periksa minimum peserta.',
        'Untuk fishing, periksa jam, jenis pemancingan, aturan, dan kuota joran.',
        'Periksa storage/logs/laravel.log untuk detail teknis.',
    ])
    add_heading(doc, '15.2 Pembayaran tidak menjadi paid', 2)
    add_numbered(doc, [
        'Periksa Server Key dan Client Key sesuai environment.',
        'Pastikan MIDTRANS_IS_PRODUCTION sesuai sandbox atau production.',
        'Pastikan endpoint webhook mengarah ke /midtrans/notification.',
        'Periksa signature dan gross amount pada log.',
        'Periksa record pada tabel pembayaran.',
        'Untuk lokal, gunakan endpoint check-payment hanya di environment local.',
    ])
    add_heading(doc, '15.3 Email tidak terkirim', 2)
    add_numbered(doc, [
        'Pastikan email_notification aktif di Settings.',
        'Periksa MAIL_HOST, MAIL_PORT, username, password, encryption, dan from address.',
        'Pastikan alamat email pelanggan valid.',
        'Periksa limit email harian.',
        'Periksa folder spam dan storage/logs/laravel.log.',
    ])
    add_heading(doc, '15.4 Foto atau perubahan paket tidak terlihat', 2)
    add_bullets(doc, [
        'Pastikan format dan ukuran file sesuai batas upload.',
        'Pastikan paket disimpan dalam status aktif bila ingin tampil publik.',
        'Refresh dengan Ctrl+F5 atau gunakan private window.',
        'Bersihkan cache aplikasi bila diperlukan.',
    ])
    add_heading(doc, '15.5 Website menampilkan error 500', 2)
    add_bullets(doc, [
        'Pada local, cek detail exception dan laravel.log.',
        'Pada production, jangan mengaktifkan APP_DEBUG=true.',
        'Periksa koneksi database, migration, storage permission, dan konfigurasi environment.',
        'Setelah perbaikan konfigurasi, bersihkan cache config dan route sesuai SOP deployment.',
    ])

    add_heading(doc, '16. Keamanan dan Praktik Operasional', 1)
    add_bullets(doc, [
        'Jangan membagikan file .env, APP_KEY, Server Key Midtrans, password SMTP, atau token lain.',
        'Gunakan APP_DEBUG=false pada production.',
        'Ganti password admin default setelah instalasi.',
        'Gunakan HTTPS pada website production dan endpoint webhook.',
        'Batasi akses panel admin hanya untuk operator yang diperlukan.',
        'Jangan menandai pembayaran paid berdasarkan bukti yang belum diverifikasi.',
        'Backup database secara berkala dan uji pemulihan backup.',
        'Periksa log setelah perubahan Midtrans, SMTP, migration, atau scheduler.',
        'Gunakan mode simulation hanya untuk pengujian lokal, bukan untuk operasional nyata.',
        'Jika credential pernah tersebar, segera rotasi credential di provider terkait.',
    ])
    add_note(doc, 'Data pelanggan', 'Kode booking berfungsi sebagai identifier akses halaman konfirmasi/e-ticket. Perlakukan kode booking sebagai data sensitif dan jangan menyebarkannya di kanal publik.')

    add_heading(doc, '17. Checklist Operasional Admin', 1)
    add_heading(doc, '17.1 Checklist harian', 2)
    checklist = [
        'Login ke dashboard dan periksa booking pending.',
        'Periksa pembayaran baru dan status Midtrans.',
        'Periksa booking untuk kunjungan hari ini dan besok.',
        'Periksa kuota paket dan fishing.',
        'Periksa notifikasi admin yang belum dibaca.',
        'Periksa email gagal atau pesan pelanggan yang perlu ditindaklanjuti.',
        'Pastikan perubahan status booking sudah sesuai bukti pembayaran.',
    ]
    for item in checklist:
        paragraph = doc.add_paragraph(style='List Bullet')
        paragraph.add_run('[ ] ' + item)
    add_heading(doc, '17.2 Checklist mingguan', 2)
    weekly = [
        'Export laporan booking dan pendapatan.',
        'Periksa paket aktif, harga, kuota, dan foto.',
        'Periksa log aplikasi dan error pembayaran/email.',
        'Pastikan scheduler pengingat dan review berjalan.',
        'Verifikasi backup database terbaru.',
    ]
    for item in weekly:
        paragraph = doc.add_paragraph(style='List Bullet')
        paragraph.add_run('[ ] ' + item)
    add_heading(doc, '17.3 Checklist sebelum production', 2)
    production = [
        'APP_ENV dan APP_DEBUG sudah benar.',
        'APP_KEY tersedia dan tidak dibagikan.',
        'Midtrans production/sandbox sudah sesuai tujuan.',
        'Webhook Midtrans sudah diarahkan ke URL HTTPS yang benar.',
        'SMTP berhasil diuji.',
        'Password admin default sudah diganti.',
        'Backup dan monitoring sudah disiapkan.',
    ]
    for item in production:
        paragraph = doc.add_paragraph(style='List Bullet')
        paragraph.add_run('[ ] ' + item)

    add_heading(doc, '18. Ringkasan Endpoint dan Komponen Sistem', 1)
    add_heading(doc, '18.1 Endpoint publik utama', 2)
    add_table(doc, ['Method', 'Endpoint', 'Fungsi'], [
        ('GET', '/', 'Halaman utama.'),
        ('GET', '/destinasi/{slug}', 'Detail destinasi.'),
        ('GET', '/paket/{category}', 'Kategori paket.'),
        ('GET', '/paket/{slug}', 'Detail paket.'),
        ('GET', '/wisata-edukasi', 'Halaman wisata edukasi.'),
        ('GET', '/wisata-edukasi/sekolah-mitra', 'Daftar sekolah mitra.'),
        ('GET', '/kontak', 'Halaman kontak.'),
        ('POST', '/api/booking/store', 'Booking reguler.'),
        ('POST', '/api/booking/fishing', 'Booking fishing.'),
        ('GET', '/api/booking/status/{kode}', 'Cek status booking.'),
        ('POST', '/midtrans/notification', 'Webhook pembayaran.'),
        ('GET', '/booking/confirmation/{kode}', 'Konfirmasi booking.'),
        ('GET', '/booking/e-ticket/{kode}', 'Download e-ticket.'),
    ], widths=[0.75, 2.75, 3.2])
    add_heading(doc, '18.2 Komponen backend penting', 2)
    add_table(doc, ['Komponen', 'Peran'], [
        ('BookingController', 'Menerima request booking, status, webhook, konfirmasi, dan e-ticket.'),
        ('BookingCreationService', 'Membuat booking, jadwal, kuota, order, dan pembayaran.'),
        ('BookingPricingService', 'Menghitung total harga booking.'),
        ('PaymentStatusService', 'Memetakan status pembayaran dan menjalankan efek sukses.'),
        ('MidtransConfigService', 'Mengatur konfigurasi Midtrans.'),
        ('BookingEmailNotificationService', 'Mengirim email booking, sukses, dan pembatalan.'),
        ('NotificationService', 'Membuat notifikasi admin.'),
        ('ETicketService', 'Membuat dan mengunduh PDF e-ticket.'),
        ('SEOService', 'Menyediakan metadata SEO dan canonical URL.'),
    ], widths=[2.4, 4.3])
    add_heading(doc, 'Penutup', 1)
    add_para(doc, 'Website Godong Ijo menggabungkan informasi destinasi dengan alur booking dan operasional admin. Untuk penggunaan harian, fokus utama admin adalah memantau booking pending, memverifikasi pembayaran, menjaga kuota, menindaklanjuti email, dan memastikan konten paket tetap akurat.')
    add_note(doc, 'Dokumen ini', 'Manual book ini dibuat berdasarkan struktur route, controller, service, view, konfigurasi, dan test yang tersedia pada workspace. Perbarui dokumen bila ada perubahan besar pada alur booking, pembayaran, menu admin, atau konfigurasi deployment.')

    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    build_document()
